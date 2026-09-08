"""Estrazione delle icone dei campioni incorporate nel file Excel.

Le icone sono immagini "oneCellAnchor" nel drawing layer di xl/drawings/,
una per riga della colonna "Icona" - openpyxl non le espone tramite la sua
API immagini semplificata, quindi si legge l'XML del drawing direttamente.
"""

import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl

from driftdraft.data import DEFAULT_DATA_PATH, NAME_COLUMN, SHEET_NAME

_NS = {
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}
_R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"

DEFAULT_ICONS_DIR = Path(__file__).resolve().parent.parent / "assets" / "icons"


def _resolve_zip_path(base_dir: str, relative_target: str) -> str:
    # Il Target di una relationship puo' essere RELATIVO alla cartella del
    # part che lo referenzia (es. "../media/image1.png" - il file xlsx
    # originale) oppure ASSOLUTO rispetto alla radice del pacchetto (es.
    # "/xl/media/image1.png", slash iniziale - convenzione diversa ma
    # ugualmente valida per lo standard OPC/OOXML, quella che openpyxl
    # scrive quando genera un file da zero). BUG REALE TROVATO 2026-08-20
    # (ricostruzione xlsx via openpyxl, vedi note di progetto): la vecchia
    # versione assumeva sempre relativo, un target assoluto produceva un
    # percorso doppiato e sbagliato (es. "xl/drawings/xl/media/image1.png").
    if relative_target.startswith("/"):
        return relative_target.lstrip("/")
    parts = base_dir.split("/") if base_dir else []
    for part in relative_target.split("/"):
        if part == "..":
            parts.pop()
        elif part not in (".", ""):
            parts.append(part)
    return "/".join(parts)


def extract_champion_icons(path: Path = DEFAULT_DATA_PATH) -> dict[str, bytes]:
    """Ritorna {nome_campione: bytes_png} leggendo le icone incorporate."""

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[SHEET_NAME]
    header = [c.value for c in ws[1]]
    name_col = header.index(NAME_COLUMN) + 1  # 1-indexed per openpyxl

    icons: dict[str, bytes] = {}

    with zipfile.ZipFile(path) as z:
        drawing_parts = [
            n for n in z.namelist() if re.fullmatch(r"xl/drawings/drawing\d+\.xml", n)
        ]
        if not drawing_parts:
            return icons
        drawing_part = drawing_parts[0]
        drawing_dir = drawing_part.rsplit("/", 1)[0]
        rels_part = f"{drawing_dir}/_rels/{drawing_part.rsplit('/', 1)[1]}.rels"

        rid_to_target = {}
        if rels_part in z.namelist():
            rels_xml = ET.fromstring(z.read(rels_part))
            for rel in rels_xml.findall("rel:Relationship", _NS):
                rid_to_target[rel.get("Id")] = rel.get("Target")

        drawing_xml = ET.fromstring(z.read(drawing_part))
        anchors = drawing_xml.findall("xdr:oneCellAnchor", _NS) + drawing_xml.findall(
            "xdr:twoCellAnchor", _NS
        )

        for anchor in anchors:
            from_el = anchor.find("xdr:from", _NS)
            blip = anchor.find(".//a:blip", _NS)
            if from_el is None or blip is None:
                continue

            row0 = from_el.find("xdr:row", _NS)
            if row0 is None:
                continue
            excel_row = int(row0.text) + 1  # xdr:row e' 0-indexed

            rid = blip.get(_R_EMBED)
            target = rid_to_target.get(rid)
            if not target:
                continue

            champ_name = ws.cell(row=excel_row, column=name_col).value
            if not champ_name:
                continue

            zip_path = _resolve_zip_path(drawing_dir, target)
            icons[champ_name] = z.read(zip_path)

    return icons


def sync_icons_to_disk(
    path: Path = DEFAULT_DATA_PATH, out_dir: Path = DEFAULT_ICONS_DIR, force: bool = False
) -> dict[str, Path]:
    """Estrae le icone e le scrive come file .png in out_dir, uno per campione
    - MA SOLO per nomi che non hanno gia' un file (a meno di force=True).

    Le icone incorporate nell'xlsx sono a bassa risoluzione (32x32 - Excel le
    comprime alla dimensione della cella quando le inserisci, trovato in un
    caso reale 2026-08-17): da quella data out_dir e' popolata a mano
    dall'utente con icone ad alta risoluzione (stesso identico stile a
    bordi colorati, sorgente originale prima della compressione di Excel),
    NON piu' da questa funzione. Sovrascrivere per errore un'icona buona con
    la versione a bassa risoluzione incorporata nell'xlsx sarebbe un
    downgrade silenzioso - default sicuro: salta chi esiste gia', riempie
    solo i nomi davvero mancanti (es. un campione nuovo aggiunto
    all'xlsx dopo). force=True per il vecchio comportamento (sovrascrive
    tutto), da usare solo se si e' sicuri che out_dir non contenga piu'
    nulla di meglio della versione xlsx."""

    out_dir.mkdir(parents=True, exist_ok=True)
    icons = extract_champion_icons(path)

    written = {}
    skipped = []
    for name, data in icons.items():
        file_path = out_dir / f"{name}.png"
        if file_path.exists() and not force:
            skipped.append(name)
            continue
        file_path.write_bytes(data)
        written[name] = file_path

    if skipped:
        print(f"Saltati {len(skipped)} campioni con icona gia' presente (usa force=True per sovrascrivere).")

    return written
