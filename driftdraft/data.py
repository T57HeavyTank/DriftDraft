import re
from dataclasses import dataclass, field, replace
from pathlib import Path

import openpyxl

from driftdraft.champion_overrides import load_overrides
from driftdraft.paths import get_app_dir

# Prima puntava al Desktop personale dell'utente (C:\Users\Tank\Desktop\...),
# inutilizzabile su qualunque altra macchina (copia sorgente originale li'
# invariata). Ora vive accanto all'eseguibile (get_app_dir(), MAI
# get_bundle_dir()) - deliberatamente FUORI da _internal/, cosi' chi riceve
# il pacchetto lo trova subito e puo' modificarlo (nuovi campioni, patch)
# senza scovarlo dentro la cartella interna di PyInstaller - e soprattutto,
# le modifiche vengono DAVVERO rilette dall'app al prossimo avvio.
DEFAULT_DATA_PATH = get_app_dir() / "data" / "champions.xlsx"
SHEET_NAME = "Campioni LoL"

NAME_COLUMN = "Campione"

COMP_COLUMNS = [
    "TeamFight \\ WomboCombo",
    "Split",
    "Pick",
    "Poke \\ Siege",
    "Proteggi il presidente",
]

ROLE_COLUMNS = ["Top", "Jungle", "Mid", "Bot", "Support"]

TAG_COLUMNS = [
    "Danno burst singolo",
    "Danno DPS singolo",
    "Danno burst AoE",
    "Danno DPS AoE",
    "Hard Engage singolo",
    "Hard Engage AoE",
    "Hard CC Singolo",
    "Hard CC AoE",
    "Mobilità",
    "Disingaggio",
    "Protezione \\ peeling",
    "Utility generica",
    "Long Range",
    "Waveclear",
    "Side forte",
]


_SLUG_SPECIAL = {"Nunu & Willump": "nunu", "Renata Glasc": "renata", "Wukong": "wukong"}


def slugify_champion_name(name: str) -> str:
    """Slug usato sia da op.gg che da lolalytics per questo campione - stessa
    convenzione su entrambi i siti (verificato), condivisa qui per evitare
    di mantenere due copie della stessa logica/eccezioni."""
    if name in _SLUG_SPECIAL:
        return _SLUG_SPECIAL[name]
    return re.sub(r"[^a-z0-9]", "", name.lower())


@dataclass(frozen=True)
class Champion:
    name: str
    comps: frozenset[str] = field(default_factory=frozenset)
    tags: frozenset[str] = field(default_factory=frozenset)
    roles: frozenset[str] = field(default_factory=frozenset)
    riot_id: int | None = None


_champions_cache: list[Champion] | None = None
_champions_cache_key: tuple | None = None


def load_champions(path: Path = DEFAULT_DATA_PATH) -> list[Champion]:
    global _champions_cache, _champions_cache_key

    st = path.stat()
    key = (str(path), st.st_mtime_ns, st.st_size)
    if _champions_cache is not None and _champions_cache_key == key:
        return _champions_cache

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[SHEET_NAME]

    header_row = [cell.value for cell in ws[1]]
    col_index = {name: idx for idx, name in enumerate(header_row)}

    champions = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = row[col_index[NAME_COLUMN]]
        if not name:
            continue

        comps = frozenset(
            c for c in COMP_COLUMNS if row[col_index[c]] not in (None, "")
        )
        tags = frozenset(
            t for t in TAG_COLUMNS if row[col_index[t]] not in (None, "")
        )
        roles = frozenset(
            r for r in ROLE_COLUMNS if row[col_index[r]] not in (None, "")
        )
        riot_id_value = row[col_index["ID"]] if "ID" in col_index else None
        riot_id = int(riot_id_value) if riot_id_value not in (None, "") else None
        champions.append(
            Champion(name=name, comps=comps, tags=tags, roles=roles, riot_id=riot_id)
        )

    # Override locali per-campione fatti dal coach in app (richiesto
    # esplicitamente 2026-08-27, vedi champion_overrides.py per il perche'
    # e' un file separato invece di riscrivere l'xlsx) - applicati QUI,
    # dentro load_champions() stessa, cosi' OGNI chiamante (server.py,
    # training_bot.py, comps.py, drafter_live.py...) li vede automaticamente
    # senza dover ricordarsi di applicarli a parte in ciascun punto.
    overrides = load_overrides()
    if overrides:
        champions = [
            replace(
                c,
                comps=frozenset(overrides[c.name]["comps"]) if "comps" in overrides.get(c.name, {}) else c.comps,
                tags=frozenset(overrides[c.name]["tags"]) if "tags" in overrides.get(c.name, {}) else c.tags,
                roles=frozenset(overrides[c.name]["roles"]) if "roles" in overrides.get(c.name, {}) else c.roles,
            )
            if c.name in overrides
            else c
            for c in champions
        ]

    _champions_cache = champions
    _champions_cache_key = key
    return champions
